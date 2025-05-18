from .text_model import anonymize_text
from .image_model import anonymize_image

from docx import Document


def anonymize_data(user_input_type, data):
    if user_input_type == 'text':
        return anonymize_text(data)
    
    elif user_input_type == "txt" or user_input_type == "md" or user_input_type == "csv":
        with open(data, "r") as file:
            text = file.read()
        text = anonymize_text(text)
        with open(data, "w") as file:
            file.write(text)
        return data
    
    elif user_input_type == "docx":
        doc = Document(data)

        for para in doc.paragraphs:
            para.text = anonymize_text(para.text)

        doc.save(data)

    elif user_input_type == "xlsx":
        wb = load_workbook(data)

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = anonymize_text(cell.value)

        wb.save(data)
    
    elif user_input_type in ['png', 'jpg', 'jpeg', 'bmp', 'webp', 'ppm', 'tiff', 'tif', 'pgm', 'pbm']:
        anonymize_image(data, data)

    elif user_input_type == 'file_info':
        filename = data.get('filename', 'неизвестный файл')
        return f"[Содержимое файла '{filename}' было анонимизировано]"

    return "Неизвестный тип данных для анонимизации (внутренняя ошибка)."
